"""
Offline test: which way of cutting documents into pieces ranks the answer best?

    uv run python evaluation/compare_splits.py

Cuts the 12 test documents (plans/documents/test-files) six ways, embeds the
pieces with the app's own embedding call, and for each of the 10 evaluation
questions searches only the documents that person may see, using the search
phrase the agent really used (results/answers.json). Nothing in the app or the
database changes; it needs only the Gemini key (~15 embedding requests).

    A today          1,000-character cuts (slides 500), as ingestion does now
    B by heading     a new piece at every heading; long sections cut at 1,000
    C B + titles     B, with "Document title > Section" at the top of each piece
    D C + merging    C, but small neighbouring sections joined up to ~400 characters
    F C + merging    like D, but only sections under ~200 characters are joined
    E parent-child   search C's pieces, but hand the agent each match's parent: a
                     block of whole neighbouring sections up to ~1,000 characters

Scores come from a marker phrase that only the answer piece contains (no AI
judge), so they are stand-ins for the DeepEval measures:
    answer first   the answer piece is ranked 1st (Context Precision)
    precision      DeepEval's ranking formula with marker-based relevance
    useful share   share of returned pieces that hold the answer (Context Relevancy)
    chars sent     text handed to the agent per question (less = less noise)
    found          the answer is in the returned pieces at all (Context Recall)

Writes results/splitting_comparison.md.
"""
import asyncio
import io
import json
import math
import re
import sys
from pathlib import Path

EVAL_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(EVAL_DIR.parent))

from langchain_text_splitters import RecursiveCharacterTextSplitter  # noqa: E402

from app.services.embeddings import _embed, embed_documents  # noqa: E402
from app.services.ingestion import _chunk_text, _extract_text  # noqa: E402

DOCS_DIR = EVAL_DIR.parent.parent / "plans" / "documents" / "test-files"
ANSWERS = EVAL_DIR / "results" / "answers.json"
REPORT = EVAL_DIR / "results" / "splitting_comparison.md"
CACHE = EVAL_DIR / ".cache" / "embeddings.json"   # git-ignored; reruns only embed new text


async def cached_embed(texts: list[str], task_type: str) -> list[list[float]]:
    """Embed only texts not embedded before (the free tier limits text per minute)."""
    cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}
    keys = [f"{task_type}|{t}" for t in texts]
    missing = list(dict.fromkeys(k for k in keys if k not in cache))
    if missing:
        vectors = await (_embed([k.split("|", 1)[1] for k in missing], task_type) if task_type == "RETRIEVAL_QUERY"
                         else embed_documents([k.split("|", 1)[1] for k in missing]))
        cache.update(zip(missing, vectors))
        CACHE.parent.mkdir(exist_ok=True)
        CACHE.write_text(json.dumps(cache), encoding="utf-8")
    return [cache[k] for k in keys]

MAX_PIECE = 1000
MIN_MERGED = 400
_LONG_SECTION = RecursiveCharacterTextSplitter(
    chunk_size=MAX_PIECE, chunk_overlap=200, separators=["\n\n", "\n", ".", " ", ""]
)

# Who sees which document (test-files README §2; the employee account holds
# TEST-Employee on TEST-Alpha, so it still sees the handbook and the charter).
ORG_WIDE = ["org-code-of-conduct", "org-travel-and-expense-policy", "org-it-security-policy",
            "org-holiday-calendar-2026"]
VISIBLE = {
    "TEST-Employee": ORG_WIDE + ["employee-handbook", "team-alpha-project-charter"],
    "TEST-Manager": ORG_WIDE + ["manager-playbook", "manager-performance-reviews",
                                "consultant-engagement-guide", "employee-handbook"],
    "TEST-HR": ORG_WIDE + ["hr-recruitment-and-onboarding", "hr-compensation-and-payroll"],
}

# A phrase only the answer piece contains, per question id (any one counts)
MARKERS = {
    1: ["2 floating holidays"],
    2: ["economy class", "25,000"],
    3: ["within 1 hour"],
    4: ["Never share your password", "IT resets the account"],
    5: ["18 days"],
    6: ["10:15"],
    7: ["1 working day"],
    8: ["reviews in a row"],
    9: ["L3 Senior Engineer"],
    10: ["60 days"],
}

_NUMBERED_HEADING = re.compile(r"^\s*\d{1,2}\.\s+[A-Za-z][^.:]{0,60}$")


# ── Reading documents ─────────────────────────────────────────────────────────

def read_with_headings(path: Path) -> str:
    """Like ingestion's _extract_text, but a docx heading keeps a '## ' marker."""
    if path.suffix != ".docx":
        return _extract_text(path.read_bytes(), path.suffix[1:])
    from docx import Document as DocxDocument
    doc = DocxDocument(io.BytesIO(path.read_bytes()))
    lines = [("## " if p.style.name.startswith("Heading") else "") + p.text
             for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def is_heading(line: str, file_type: str) -> bool:
    if file_type in ("md", "docx"):
        return line.lstrip().startswith("#")
    if file_type in ("txt", "pdf"):           # only numbered lines like "1. Passwords and sign-in"
        return bool(_NUMBERED_HEADING.match(line))
    return False


def sections(text: str, file_type: str) -> list[tuple[str, str]]:
    """(heading, text) per section. Slides are one section each."""
    if file_type == "pptx":
        return [(block.split("]")[0].strip("["), block) for block in text.split("\n\n") if block.strip()]
    result: list[tuple[str, list[str]]] = [("", [])]
    for line in text.splitlines():
        if is_heading(line, file_type):
            result.append((line.strip("# ").strip(), [line]))
        else:
            result[-1][1].append(line)
    return [(h, "\n".join(lines).strip()) for h, lines in result if "\n".join(lines).strip()]


def title_of(text: str) -> str:
    first = next(line for line in text.splitlines() if line.strip())
    return first.strip("# ").replace("(Test Document)", "").replace("(TEST DOCUMENT)", "").strip(" -")


# ── The six ways of cutting ──────────────────────────────────────────────────

def cut_today(text: str, file_type: str, title: str) -> list[str]:
    return _chunk_text(text, file_type)


def cut_by_heading(text: str, file_type: str, title: str, with_title=False, merge=False,
                   min_size=MIN_MERGED) -> list[str]:
    parts = sections(text, file_type)
    if merge:
        merged: list[tuple[str, str]] = []
        for heading, body in parts:
            if merged and len(merged[-1][1]) < min_size and len(merged[-1][1]) + len(body) <= MAX_PIECE:
                last_heading, last_body = merged[-1]
                merged[-1] = (" · ".join(h for h in (last_heading, heading) if h), last_body + "\n\n" + body)
            else:
                merged.append((heading, body))
        parts = merged
    pieces = []
    for heading, body in parts:
        for piece in (_LONG_SECTION.split_text(body) if len(body) > MAX_PIECE else [body]):
            if with_title:
                piece = f"{title} > {heading}\n{piece}" if heading else f"{title}\n{piece}"
            pieces.append(piece)
    return pieces


def cut_parent_child(text: str, file_type: str, title: str) -> list[tuple[str, str]]:
    """
    (child, parent) pairs. Children are C's pieces (one per section, with the
    title) and are what gets embedded and searched. Parents are blocks of whole
    neighbouring sections up to ~1,000 characters, never cut inside a section;
    the agent is given the parent of each matching child.
    """
    blocks: list[list[tuple[str, str]]] = []
    for heading, body in sections(text, file_type):
        if blocks and sum(len(b) for _, b in blocks[-1]) + len(body) <= MAX_PIECE:
            blocks[-1].append((heading, body))
        else:
            blocks.append([(heading, body)])
    pairs = []
    for block in blocks:
        parent = f"{title}\n" + "\n\n".join(body for _, body in block)
        for heading, body in block:
            for piece in (_LONG_SECTION.split_text(body) if len(body) > MAX_PIECE else [body]):
                pairs.append((f"{title} > {heading}\n{piece}" if heading else f"{title}\n{piece}", parent))
    return pairs


STRATEGIES = {
    "A today (1,000 cuts)": cut_today,
    "B by heading": lambda t, f, ti: cut_by_heading(t, f, ti),
    "C by heading + titles": lambda t, f, ti: cut_by_heading(t, f, ti, with_title=True),
    "D by heading + titles + merging": lambda t, f, ti: cut_by_heading(t, f, ti, with_title=True, merge=True),
    "E parent-child (search C, return section blocks)": cut_parent_child,
    "F by heading + titles + merging under 200": lambda t, f, ti: cut_by_heading(t, f, ti, with_title=True, merge=True, min_size=200),
}


# ── Scoring ───────────────────────────────────────────────────────────────────

def cosine(a: list[float], b: list[float]) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    return dot / (math.sqrt(sum(x * x for x in a)) * math.sqrt(sum(y * y for y in b)))


def precision(relevant: list[bool]) -> float:
    """DeepEval's Contextual Precision formula, with marker-based relevance."""
    hits, total = 0, 0.0
    for position, is_relevant in enumerate(relevant, start=1):
        if is_relevant:
            hits += 1
            total += hits / position
    return total / hits if hits else 0.0


def evaluate(pieces: list[dict], query_vectors: dict[int, list[float]], answers: list[dict], k: int) -> dict:
    rows = []
    for answer in answers:
        visible = set(VISIBLE[answer["ask_as"]])
        matches = sorted(
            (p for p in pieces if p["doc"] in visible),
            key=lambda p: cosine(query_vectors[answer["id"]], p["vector"]),
            reverse=True,
        )
        # What the agent gets: each match's returned text (its parent for E), no repeats
        returned: list[str] = []
        for p in matches:
            if p["returned"] not in returned:
                returned.append(p["returned"])
            if len(returned) == k:
                break
        relevant = [any(m in text for m in MARKERS[answer["id"]]) for text in returned]
        rows.append({
            "id": answer["id"],
            "first": relevant[0],
            "rank": relevant.index(True) + 1 if any(relevant) else None,
            "precision": precision(relevant),
            "useful": sum(relevant) / len(relevant),
            "chars": sum(len(text) for text in returned),
            "found": any(relevant),
        })
    n = len(rows)
    return {
        "rows": rows,
        "first": sum(r["first"] for r in rows),
        "precision": sum(r["precision"] for r in rows) / n,
        "useful": sum(r["useful"] for r in rows) / n,
        "chars": sum(r["chars"] for r in rows) // n,
        "found": sum(r["found"] for r in rows),
    }


async def main() -> None:
    answers = json.loads(ANSWERS.read_text(encoding="utf-8"))
    files = [f for f in sorted(DOCS_DIR.iterdir()) if f.suffix in (".md", ".txt", ".pdf", ".docx", ".pptx")
             and f.name != "README.md"]

    print("Embedding the 10 search phrases ...")
    phrases = [a["searches"][0] for a in answers]
    query_vectors = dict(zip((a["id"] for a in answers), await cached_embed(phrases, "RETRIEVAL_QUERY")))

    results = {}
    for name, cut in STRATEGIES.items():
        pieces = []
        for f in files:
            file_type = f.suffix[1:]
            text = read_with_headings(f) if name[0] != "A" else _extract_text(f.read_bytes(), file_type)
            for piece in cut(text, file_type, title_of(text)):
                child, parent = piece if isinstance(piece, tuple) else (piece, piece)
                pieces.append({"doc": f.stem, "text": child, "returned": parent})
        print(f"{name}: {len(pieces)} pieces, embedding ...")
        for piece, vector in zip(pieces, await cached_embed([p["text"] for p in pieces], "RETRIEVAL_DOCUMENT")):
            piece["vector"] = vector
        results[name] = {
            "pieces": len(pieces),
            "avg_size": sum(len(p["text"]) for p in pieces) // len(pieces),
            "k6": evaluate(pieces, query_vectors, answers, 6),
            "k4": evaluate(pieces, query_vectors, answers, 4),
            "k3": evaluate(pieces, query_vectors, answers, 3),
        }

    write_report(results, answers)
    print(f"Report: {REPORT.relative_to(EVAL_DIR.parent)}")


def write_report(results: dict, answers: list[dict]) -> None:
    lines = [
        "# Splitting comparison (offline)",
        "",
        "The 12 test documents cut six ways and searched with the agent's real search phrases, "
        "each question limited to the documents that person may see. Relevance comes from a marker phrase "
        "that only the answer piece contains (no AI judge), so these are stand-ins for the DeepEval measures. "
        "Made by `evaluation/compare_splits.py`; nothing in the app changed.",
        "",
        "| Way of cutting | Pieces | Avg size | Returned | Answer 1st | Precision | Useful share | Chars sent | Found |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for name, r in results.items():
        for k in (6, 4, 3):
            m = r[f"k{k}"]
            lines.append(
                f"| {name} | {r['pieces']} | {r['avg_size']} | {k} | {m['first']}/10 | {m['precision']:.2f} | "
                f"{m['useful']:.2f} | {m['chars']} | {m['found']}/10 |"
            )
    lines += [
        "",
        "- **Answer 1st / Precision**: stand-ins for Context Precision (the live run scored 0.75).",
        "- **Useful share / Chars sent**: stand-ins for Context Relevancy (live 0.13): more useful pieces and "
        "less text sent means less noise.",
        "- **Found**: stand-in for Context Recall (live 1.00). Must stay 10/10.",
        "- **Pieces / Avg size** count what is embedded (for E, the children); **Chars sent** counts what "
        "the agent receives (for E, the parents).",
        "",
        "## Rank of the answer piece per question (6 returned)",
        "",
        "| # | Question | " + " | ".join(n.split(" ")[0] for n in results) + " |",
        "|---|---|" + "---|" * len(results),
    ]
    for i, answer in enumerate(answers):
        ranks = [results[n]["k6"]["rows"][i]["rank"] for n in results]
        lines.append(f"| {answer['id']} | {answer['question']} | "
                     + " | ".join(str(r) if r else "not found" for r in ranks) + " |")
    lines += ["", "Strategies: " + "; ".join(results), ""]
    REPORT.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    asyncio.run(main())
