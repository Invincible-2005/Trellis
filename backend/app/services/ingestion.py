"""
Document Ingestion Pipeline  (Ticket 11)
=========================================
Fetch → Extract → Chunk → Embed → Bulk-insert DocumentChunks

Called exclusively as a FastAPI BackgroundTask from POST /api/v1/documents,
as `background_tasks.add_task(run_ingestion, document.id)`. It opens its own
DB session — never pass the request's session into a background task.

Security invariants:
    - Scope columns (org_id, scope_type, role_id, team_id) are copied from the
      parent Document row — they are set by the authenticated upload endpoint,
      never derived from file contents. The LLM cannot influence them.
    - Only UploadThing CDN URLs are fetched (`is_allowed_file_url`), over
      HTTPS, without following redirects, and capped at
      INGESTION_MAX_FILE_BYTES. This prevents SSRF against internal hosts.

Consistency invariant:
    Chunks and `status='ready'` are committed in ONE transaction. A document
    is either ready with all its chunks, or failed with none, so retrieval
    can never see chunks of a failed or half-ingested document.

Chunking decisions  (Ticket 11):
    Standard docs  (.pdf, .docx, .txt, .md) : 1000 chars / 200 overlap
    Presentations  (.pptx)                   : 500 chars  / 100 overlap

Embedding  (Ticket 12 D14):
    gemini-embedding-001 → 1024-dim vectors, via app/services/embeddings.py
    (embed_documents here; search_documents uses embed_query from the same module).
"""

import io
import asyncio
import logging
import uuid
from datetime import datetime, timedelta
from typing import Optional
from urllib.parse import urlsplit

import httpx
from sqlalchemy import delete, update
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.core.config import settings
from app.core.database import async_session
from app.models.document import Document, DocumentChunk
from app.services.embeddings import EmbeddingError, embed_documents

logger = logging.getLogger(__name__)


class IngestionError(Exception):
    """Expected failure. Its message is admin-safe and stored in documents.error_message."""


# ─── Splitter instances — created once at module load, reused per call ────────

_STANDARD_SPLITTER = RecursiveCharacterTextSplitter(
    chunk_size=1000,
    chunk_overlap=200,
    separators=["\n\n", "\n", ".", " ", ""],
)

_PPTX_SPLITTER = RecursiveCharacterTextSplitter(
    chunk_size=500,
    chunk_overlap=100,
    separators=["\n\n", "\n", " • ", ". ", " ", ""],
)


# ─── Step 1: Fetch raw bytes from UploadThing CDN ─────────────────────────────

def is_allowed_file_url(file_url: str) -> bool:
    """
    True only for https URLs on an allowed UploadThing host (SSRF guard).
    Entries in INGESTION_ALLOWED_HOSTS starting with "." match subdomains.
    Also used by the upload endpoint to reject bad URLs before a row is created.
    """
    try:
        parts = urlsplit(file_url)
        port = parts.port
    except ValueError:
        return False
    host = (parts.hostname or "").lower()
    if parts.scheme != "https" or not host or port is not None:
        return False
    for allowed in settings.INGESTION_ALLOWED_HOSTS:
        allowed = allowed.lower()
        if allowed.startswith("."):
            if host.endswith(allowed) and len(host) > len(allowed):
                return True
        elif host == allowed:
            return True
    return False


async def _fetch_file(file_url: str) -> bytes:
    """
    Downloads raw file bytes from the UploadThing CDN URL.
    Streams the body and aborts once INGESTION_MAX_FILE_BYTES is exceeded.
    Times out after 30 s to prevent hung background tasks.
    """
    if not is_allowed_file_url(file_url):
        raise IngestionError("File URL is not an allowed UploadThing URL.")

    max_bytes = settings.INGESTION_MAX_FILE_BYTES
    too_large = IngestionError(f"File exceeds the {max_bytes // (1024 * 1024)} MB limit.")
    try:
        async with httpx.AsyncClient(timeout=30.0, follow_redirects=False) as client:
            async with client.stream("GET", file_url) as response:
                if response.is_redirect:
                    raise IngestionError("File URL redirected; only direct UploadThing URLs are allowed.")
                if response.status_code != 200:
                    raise IngestionError(f"File could not be downloaded (HTTP {response.status_code}).")

                declared = response.headers.get("content-length")
                if declared and declared.isdigit() and int(declared) > max_bytes:
                    raise too_large

                buffer = bytearray()
                async for piece in response.aiter_bytes():
                    buffer.extend(piece)
                    if len(buffer) > max_bytes:
                        raise too_large
                return bytes(buffer)
    except httpx.HTTPError as exc:
        raise IngestionError("File could not be downloaded (network error).") from exc


# ─── Step 2: Extract plain text (file-type aware) ─────────────────────────────

def _pptx_shape_texts(shapes) -> list[str]:
    """Text from text frames, tables, and (recursively) grouped shapes."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    texts: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(_pptx_shape_texts(shape.shapes))
        elif shape.has_text_frame:
            frame_text = shape.text_frame.text.strip()
            if frame_text:
                texts.append(frame_text)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    texts.append(" | ".join(cells))
    return texts


def _extract_text(raw_bytes: bytes, file_type: str) -> str:
    """
    Converts raw file bytes to a single plain-text string.
    CPU-bound — always call via asyncio.to_thread, never directly on the event loop.

    .pptx  → Each slide becomes "[Slide N] <text>" preserving search context.
    .pdf   → pymupdf — fastest, handles multi-column layouts.
    .docx  → python-docx — paragraphs, then table rows.
    .txt / .md → UTF-8 decode, no transformation.
    """
    if file_type in ("txt", "md"):
        return raw_bytes.decode("utf-8", errors="replace")

    if file_type == "pdf":
        import pymupdf
        with pymupdf.open(stream=raw_bytes, filetype="pdf") as doc:
            pages = [page.get_text() for page in doc]
        return "\n\n".join(p for p in pages if p.strip())

    if file_type == "docx":
        from docx import Document as DocxDocument
        doc = DocxDocument(io.BytesIO(raw_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table rows so tabular policy data is indexed
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return "\n".join(paragraphs)

    if file_type == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(raw_bytes))
        slide_texts: list[str] = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = _pptx_shape_texts(slide.shapes)
            # has_notes_slide: reading slide.notes_slide would CREATE an empty one
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                note_text = slide.notes_slide.notes_text_frame.text.strip()
                if note_text:
                    texts.append(f"[Notes] {note_text}")
            if texts:
                slide_texts.append(f"[Slide {i}] " + " ".join(texts))
        return "\n\n".join(slide_texts)

    raise IngestionError(f"Unsupported file type: '{file_type}'.")


# ─── Step 3: Chunk text ───────────────────────────────────────────────────────

def _chunk_text(text: str, file_type: str) -> list[str]:
    """
    Splits extracted text into overlapping chunks.
    PPTX uses smaller chunks (500/100) because slide bullets are already short;
    larger chunks would merge unrelated slides into one embedding.
    """
    splitter = _PPTX_SPLITTER if file_type == "pptx" else _STANDARD_SPLITTER
    chunks = splitter.split_text(text)
    # Drop empty chunks that can result from sparse slides or blank pages
    return [c for c in chunks if c.strip()]


# ─── Step 4: Embed — see app/services/embeddings.py (embed_documents) ─────────


# ─── Step 5: Replace DocumentChunk rows (no commit — caller owns the transaction)

async def _replace_chunks(db, document: Document, chunks: list[str], embeddings: list[list[float]]) -> None:
    """
    Deletes any existing chunks for the document (re-runs are idempotent), then
    inserts one DocumentChunk row per chunk.

    Scope columns are DENORMALIZED from the parent Document so that the HNSW
    vector search query can filter with a simple WHERE clause instead of a JOIN:

        WHERE scope_type = 'role' AND role_id IN (SELECT id FROM visible_roles)

    This is the critical performance and security invariant from Ticket 02/11.
    """
    await db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == document.id))
    db.add_all([
        DocumentChunk(
            document_id=document.id,
            content=chunk_text,
            embedding=embedding,            # list[float] of length EMBEDDING_DIM
            chunk_index=chunk_index,
            # ── Denormalized scope (copied from parent Document) ──────────────
            org_id=document.org_id,
            scope_type=document.scope_type,
            role_id=document.role_id,       # None for team- or org-scoped docs
            team_id=document.team_id,       # None for role- or org-scoped docs
        )
        for chunk_index, (chunk_text, embedding) in enumerate(zip(chunks, embeddings))
    ])


# ─── Failure bookkeeping ──────────────────────────────────────────────────────

async def _mark_failed(db, document_id: uuid.UUID, message: str) -> None:
    """Explicit UPDATE: after a rollback the ORM instance is expired and must not be reused."""
    try:
        await db.execute(
            update(Document)
            .where(Document.id == document_id)
            .values(status="failed", chunk_count=None, error_message=message[:1000],
                    updated_at=datetime.utcnow())
        )
        await db.commit()
    except Exception:
        logger.exception("[ingestion] Could not set status=failed for %s", document_id)


async def fail_stale_ingestions(older_than_minutes: int = 15) -> int:
    """
    Marks documents stuck in 'processing' (e.g. the server restarted mid-task) as
    failed so the admin can re-upload. Called once from the FastAPI lifespan.
    Assumes a single backend instance (Ticket 12 D11).
    """
    cutoff = datetime.utcnow() - timedelta(minutes=older_than_minutes)
    async with async_session() as db:
        result = await db.execute(
            update(Document)
            .where(Document.status == "processing", Document.updated_at < cutoff)
            .values(status="failed", error_message="Interrupted by a server restart. Please re-upload.",
                    updated_at=datetime.utcnow())
        )
        await db.commit()
        return result.rowcount or 0


# ─── Orchestrator ─────────────────────────────────────────────────────────────

async def run_ingestion(document_id: uuid.UUID) -> None:
    """
    Full 5-step ingestion pipeline for a single document.

    Called by:   FastAPI BackgroundTasks (not in the HTTP request cycle)
    On success:  chunks + status='ready' + chunk_count committed together
    On failure:  status='failed', error_message set, no chunks kept
    """
    async with async_session() as db:
        document: Optional[Document] = await db.get(Document, document_id)
        if document is None:
            logger.error("[ingestion] Document %s not found — skipping.", document_id)
            return

        logger.info("[ingestion] Starting: %r (%s) id=%s", document.title, document.file_type, document_id)

        try:
            # Step 1 — Fetch
            raw_bytes = await _fetch_file(document.file_url)

            # Step 2 — Extract (CPU-bound → worker thread, keeps the event loop free)
            try:
                text = await asyncio.to_thread(_extract_text, raw_bytes, document.file_type)
            except IngestionError:
                raise
            except Exception as exc:
                raise IngestionError(
                    f"Could not read the file as .{document.file_type} — it may be corrupt."
                ) from exc
            if not text.strip():
                raise IngestionError("No extractable text — the file may be image-only or corrupt.")

            # Step 3 — Chunk
            chunks = _chunk_text(text, document.file_type)
            if not chunks:
                raise IngestionError("No extractable text — the file may be image-only or corrupt.")
            logger.info("[ingestion] %d chunks produced for %s", len(chunks), document_id)

            # Step 4 — Embed
            try:
                embeddings = await embed_documents(chunks)
            except EmbeddingError as exc:
                raise IngestionError(str(exc)) from exc

            # Step 5 — Insert + finalise in ONE transaction
            await _replace_chunks(db, document, chunks, embeddings)
            document.status = "ready"
            document.chunk_count = len(chunks)
            document.error_message = None
            await db.commit()
            logger.info("[ingestion] Complete: %s — %d chunks indexed", document_id, len(chunks))

        except IngestionError as exc:
            logger.warning("[ingestion] Failed for %s: %s", document_id, exc)
            await db.rollback()
            await _mark_failed(db, document_id, str(exc))
        except Exception:
            logger.exception("[ingestion] Unexpected failure for %s", document_id)
            await db.rollback()
            await _mark_failed(db, document_id, "Unexpected error during ingestion. Please try again.")
